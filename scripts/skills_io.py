"""Export / import skill JSON <-> XLSX for easy bulk editing.

Round-trips every skill file under ``src/data/skills/**`` to a single
workbook at ``docs/skills.xlsx`` (one ``Skills`` sheet, one row per skill).
The ``source_file`` column carries the origin path (e.g. ``player/kim``,
``enemy/realm_01``) so round-trips write each skill back to the same JSON
file it came from -- drop-in a brand-new file by setting ``source_file`` to
its desired stem (``player/<elem>`` / ``enemy/<group>``).

Usage from the repo root::

    python scripts/skills_io.py export   # JSON -> docs/skills.xlsx
    python scripts/skills_io.py import   # docs/skills.xlsx -> JSON

Pass ``--xlsx <path>`` to use a different workbook location.

Edit-friendly conventions
-------------------------
* ``dmg_scale`` is split into two columns ``dmg_atk`` and ``dmg_matk`` --
  blanks default to ``0.0``. A row always emits both keys on export so
  the JSON shape stays canonical.
* ``effects`` is a comma-separated list. Blank cell -> ``[]``.
* ``effect_chances``, ``effect_overrides``, ``summon_spec``,
  ``passive_bonus``, ``chain_skill`` are stored as compact JSON in the
  cell so nested dicts round-trip exactly.
* Boolean ``internal_only`` is written as ``TRUE`` / blank.
* ``formation_key`` is the only field where a blank cell means JSON
  ``null`` (every skill carries the key -- most are null). Other optional
  fields (``hit_count``, ``description_vi``, ``charge_bonus``, ...) drop
  out of the JSON entirely when blank.
* JSON keys are emitted in a canonical order to keep diffs minimal.
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required. Install with: pip install openpyxl")


REPO_ROOT = Path(__file__).resolve().parent.parent
SKILLS_DIR = REPO_ROOT / "src" / "data" / "skills"
DEFAULT_XLSX = REPO_ROOT / "docs" / "skills.xlsx"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F6BED")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

# Column order in the spreadsheet. ``source_file`` first so editors can
# sort/filter on it; ``key`` second so each row's identity is visible.
SHEET_COLUMNS: tuple[str, ...] = (
    "source_file",
    "key",
    "vi",
    "en",
    "realm",
    "scroll_grade",
    "category",
    "element",
    "attack_type",
    "dmg_atk",
    "dmg_matk",
    "mp_cost",
    "cooldown",
    "base_dmg",
    "hit_count",
    "true_dmg_pct",
    "charge_bonus",
    "internal_only",
    "reserved_mp_pct",
    "soul_drain_procs",
    "stat_steal_procs",
    "burst_per_mana_stack_mult",
    "burst_shield_mult",
    "auto_cast_on_stacks",
    "formation_key",
    "effects",
    "effect_chances",
    "effect_overrides",
    "summon_spec",
    "passive_bonus",
    "chain_skill",
    "description_vi",
)

# Canonical JSON key order. Mirrors what the existing skill files use so
# round-trips don't churn the diffs. Keys not listed get appended at the
# end in arrival order (so a future field doesn't crash the export).
JSON_KEY_ORDER: tuple[str, ...] = (
    "key", "vi", "en", "realm", "scroll_grade", "category", "element",
    "attack_type", "dmg_scale", "mp_cost", "cooldown", "base_dmg",
    "hit_count", "effects", "effect_chances", "formation_key",
    "effect_overrides", "reserved_mp_pct", "true_dmg_pct", "charge_bonus",
    "auto_cast_on_stacks", "internal_only", "soul_drain_procs",
    "stat_steal_procs", "burst_per_mana_stack_mult", "burst_shield_mult",
    "chain_skill", "summon_spec", "passive_bonus", "description_vi",
)

# Numeric columns — exported as numbers, imported with int/float coercion.
# ``charge_bonus`` and ``auto_cast_on_stacks`` are dict-shaped in the JSON
# (e.g. ``{"every": 1, "amount": 600}``) so they live on the JSON-cell list
# below, not here.
NUMERIC_COLUMNS: frozenset[str] = frozenset({
    "realm", "scroll_grade", "mp_cost", "cooldown", "base_dmg",
    "hit_count", "true_dmg_pct", "reserved_mp_pct",
    "soul_drain_procs", "stat_steal_procs", "burst_per_mana_stack_mult",
    "burst_shield_mult",
    "dmg_atk", "dmg_matk",
})

# Fields whose values are nested dicts/lists — serialized as compact JSON
# in a single cell so they round-trip without flattening.
JSON_CELL_COLUMNS: tuple[tuple[str, bool], ...] = (
    # (column_name, expect_dict_on_import)
    ("effect_chances", True),
    ("effect_overrides", True),
    ("summon_spec", True),
    ("passive_bonus", True),
    ("chain_skill", True),
    ("charge_bonus", True),
    ("auto_cast_on_stacks", True),
)

# Fields that are dropped when blank (i.e. omitted entirely from JSON).
# ``formation_key`` is intentionally excluded — every skill carries it.
OPTIONAL_FIELDS: frozenset[str] = frozenset({
    "scroll_grade", "hit_count", "true_dmg_pct",
    "internal_only", "reserved_mp_pct", "soul_drain_procs",
    "stat_steal_procs", "burst_per_mana_stack_mult", "burst_shield_mult",
    "description_vi",
} | {col for col, _ in JSON_CELL_COLUMNS})


# ──────────────────────────────────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────────────────────────────────

def _load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _dump_json(path: Path, data: Any) -> None:
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def _csv_join(values: list[str] | None) -> str:
    if not values:
        return ""
    return ",".join(str(v) for v in values)


def _csv_split(value: Any) -> list[str]:
    if value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    return [piece.strip() for piece in text.split(",") if piece.strip()]


def _to_bool(value: Any) -> bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"", "none", "null"}:
        return None
    if text in {"true", "1", "yes", "y", "x", "✓"}:
        return True
    if text in {"false", "0", "no", "n"}:
        return False
    raise ValueError(f"Cannot interpret {value!r} as boolean")


def _to_number(value: Any) -> float | int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    text = str(value).strip()
    if not text:
        return None
    if "." in text or "e" in text.lower():
        return float(text)
    try:
        return int(text)
    except ValueError:
        return float(text)


def _dump_cell(value: Any) -> Any:
    """Serialize a dict / list as compact JSON for cell storage."""
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _parse_json_cell(value: Any, *, expect_dict: bool = False) -> Any:
    """Inverse of ``_dump_cell``. Empty -> ``None``."""
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Cell is not valid JSON: {value!r} ({exc})")
    if expect_dict and not isinstance(parsed, dict):
        raise ValueError(f"Expected a JSON object, got {type(parsed).__name__}: {text!r}")
    return parsed


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
    ws.freeze_panes = "B2"  # freeze the source_file column too


def _autosize(ws, min_width: int = 10, max_width: int = 60) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = max(
            (len(str(c.value)) for c in column_cells if c.value is not None),
            default=0,
        )
        ws.column_dimensions[letter].width = min(max(longest + 2, min_width), max_width)


def _source_file_for(path: Path) -> str:
    """Return ``"player/kim"`` for ``src/data/skills/player/kim.json``."""
    rel = path.relative_to(SKILLS_DIR)
    return rel.with_suffix("").as_posix()


def _path_for_source(source: str) -> Path:
    """Inverse of ``_source_file_for`` -- ``"player/kim"`` -> JSON path."""
    cleaned = source.strip().strip("/")
    if not cleaned:
        raise ValueError("source_file is required on every row")
    return SKILLS_DIR / f"{cleaned}.json"


def _ordered_dict(skill: dict[str, Any]) -> dict[str, Any]:
    """Reorder a skill dict so output JSON follows ``JSON_KEY_ORDER``.

    Keys not present in the canonical order are appended at the end so a
    future field added by hand doesn't get silently dropped.
    """
    ordered: dict[str, Any] = {}
    for key in JSON_KEY_ORDER:
        if key in skill:
            ordered[key] = skill[key]
    for key, value in skill.items():
        if key not in ordered:
            ordered[key] = value
    return ordered


# ──────────────────────────────────────────────────────────────────────────
# Export: JSON → XLSX
# ──────────────────────────────────────────────────────────────────────────

def _skill_to_row(source_file: str, skill: dict[str, Any]) -> list[Any]:
    dmg = skill.get("dmg_scale") or {}
    row: dict[str, Any] = {col: "" for col in SHEET_COLUMNS}
    row["source_file"] = source_file
    row["key"] = skill.get("key", "")
    row["vi"] = skill.get("vi", "")
    row["en"] = skill.get("en", "")
    row["realm"] = skill.get("realm", "")
    row["scroll_grade"] = skill.get("scroll_grade", "")
    row["category"] = skill.get("category", "")
    row["element"] = skill.get("element", "")
    row["attack_type"] = skill.get("attack_type", "")
    row["dmg_atk"] = dmg.get("atk", 0.0)
    row["dmg_matk"] = dmg.get("matk", 0.0)
    row["mp_cost"] = skill.get("mp_cost", "")
    row["cooldown"] = skill.get("cooldown", "")
    row["base_dmg"] = skill.get("base_dmg", "")
    row["hit_count"] = skill.get("hit_count", "")
    row["true_dmg_pct"] = skill.get("true_dmg_pct", "")
    row["internal_only"] = "TRUE" if skill.get("internal_only") else ""
    row["reserved_mp_pct"] = skill.get("reserved_mp_pct", "")
    row["soul_drain_procs"] = skill.get("soul_drain_procs", "")
    row["stat_steal_procs"] = skill.get("stat_steal_procs", "")
    row["burst_per_mana_stack_mult"] = skill.get("burst_per_mana_stack_mult", "")
    row["burst_shield_mult"] = skill.get("burst_shield_mult", "")
    # formation_key carries None for non-formation skills — write empty cell
    # for null and let the importer translate empty back to None.
    fk = skill.get("formation_key", "__missing__")
    row["formation_key"] = "" if fk in (None, "__missing__") else fk
    row["effects"] = _csv_join(skill.get("effects"))
    for col, _ in JSON_CELL_COLUMNS:
        row[col] = _dump_cell(skill.get(col))
    row["description_vi"] = skill.get("description_vi", "")
    return [row[col] for col in SHEET_COLUMNS]


def export_to_xlsx(xlsx_path: Path) -> None:
    if not SKILLS_DIR.exists():
        sys.exit(f"Skills directory not found: {SKILLS_DIR}")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Skills")
    ws.append(list(SHEET_COLUMNS))

    total = 0
    files_seen = 0
    for path in sorted(SKILLS_DIR.rglob("*.json")):
        skills = _load_json(path)
        if not isinstance(skills, list):
            print(f"Warning: {path} is not a JSON list -- skipping", file=sys.stderr)
            continue
        files_seen += 1
        source_file = _source_file_for(path)
        for skill in skills:
            ws.append(_skill_to_row(source_file, skill))
            total += 1

    _style_header(ws)
    _autosize(ws)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"Exported {total} skills from {files_seen} files -> {xlsx_path}")


# ──────────────────────────────────────────────────────────────────────────
# Import: XLSX → JSON
# ──────────────────────────────────────────────────────────────────────────

def _sheet_to_dicts(ws) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out: list[dict[str, Any]] = []
    for raw in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        out.append({headers[i]: raw[i] for i in range(len(headers))})
    return out


def _row_to_skill(row: dict[str, Any]) -> dict[str, Any]:
    """Convert a sheet row back into a skill dict. Drops blank optional fields."""
    key = (row.get("key") or "").strip()
    if not key:
        raise ValueError("Row missing required 'key' column")

    skill: dict[str, Any] = {"key": key}

    for field in ("vi", "en", "category", "element", "attack_type"):
        value = row.get(field)
        skill[field] = "" if value is None else str(value).strip()

    for field in NUMERIC_COLUMNS - {"dmg_atk", "dmg_matk"}:
        num = _to_number(row.get(field))
        if num is None:
            if field not in OPTIONAL_FIELDS:
                # Required numerics get a 0 default so the JSON shape stays
                # consistent — explicit zero is better than a silent KeyError
                # downstream when the registry reads the field.
                skill[field] = 0
        else:
            skill[field] = num

    # dmg_scale always emits both keys so the canonical shape is preserved
    # even when one side is zero. Force ``float`` so 1.0 / 0.0 stay floats —
    # the existing JSON files use floats deliberately and we don't want the
    # round-trip to collapse them to ``1`` / ``0``.
    atk = _to_number(row.get("dmg_atk"))
    matk = _to_number(row.get("dmg_matk"))
    skill["dmg_scale"] = {
        "atk": float(atk) if atk is not None else 0.0,
        "matk": float(matk) if matk is not None else 0.0,
    }

    internal = _to_bool(row.get("internal_only"))
    if internal:
        skill["internal_only"] = True

    # formation_key: empty cell → None (most skills are null).
    fk_raw = row.get("formation_key")
    fk_text = "" if fk_raw is None else str(fk_raw).strip()
    skill["formation_key"] = fk_text or None

    skill["effects"] = _csv_split(row.get("effects"))

    for field, expect_dict in JSON_CELL_COLUMNS:
        try:
            parsed = _parse_json_cell(row.get(field), expect_dict=expect_dict)
        except ValueError as exc:
            raise ValueError(f"Skill {key!r} field {field!r}: {exc}")
        if parsed is not None:
            skill[field] = parsed

    desc = row.get("description_vi")
    if desc is not None and str(desc).strip():
        skill["description_vi"] = str(desc).strip()

    return _ordered_dict(skill)


def import_from_xlsx(xlsx_path: Path) -> None:
    if not xlsx_path.exists():
        sys.exit(f"Workbook not found: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "Skills" not in wb.sheetnames:
        sys.exit("Workbook is missing required sheet: 'Skills'")

    rows = _sheet_to_dicts(wb["Skills"])
    if not rows:
        sys.exit("Workbook 'Skills' sheet has no data rows")

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    seen_keys: dict[str, str] = {}
    for idx, row in enumerate(rows, start=2):  # +2 → spreadsheet line number
        source = row.get("source_file")
        source_text = "" if source is None else str(source).strip()
        if not source_text:
            raise ValueError(f"Row {idx}: 'source_file' is required")
        try:
            skill = _row_to_skill(row)
        except ValueError as exc:
            raise ValueError(f"Row {idx}: {exc}") from exc
        key = skill["key"]
        if key in seen_keys:
            raise ValueError(
                f"Duplicate skill key {key!r} appears in {seen_keys[key]!r} "
                f"and {source_text!r} (row {idx})"
            )
        seen_keys[key] = source_text
        grouped[source_text].append(skill)

    # Detect existing files that aren't represented in the workbook so the
    # user gets a clear warning rather than silent retention.
    existing = {
        _source_file_for(p) for p in SKILLS_DIR.rglob("*.json")
        if p.is_file()
    }
    written: set[str] = set()
    for source, skills in grouped.items():
        path = _path_for_source(source)
        path.parent.mkdir(parents=True, exist_ok=True)
        _dump_json(path, skills)
        written.add(source)

    stale = existing - written
    if stale:
        print(
            "Warning: existing skill files not present in workbook "
            f"(left untouched): {sorted(stale)}",
            file=sys.stderr,
        )

    total = sum(len(v) for v in grouped.values())
    print(f"Imported {total} skills into {len(grouped)} files from {xlsx_path}")


# ──────────────────────────────────────────────────────────────────────────
# CLI
# ──────────────────────────────────────────────────────────────────────────

def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "mode", choices=("export", "import"),
        help="export JSON->XLSX or import XLSX->JSON",
    )
    parser.add_argument(
        "--xlsx", type=Path, default=DEFAULT_XLSX,
        help=f"workbook path (default: {DEFAULT_XLSX.relative_to(REPO_ROOT)})",
    )
    args = parser.parse_args(argv)

    if args.mode == "export":
        export_to_xlsx(args.xlsx)
    else:
        import_from_xlsx(args.xlsx)


if __name__ == "__main__":
    main()
