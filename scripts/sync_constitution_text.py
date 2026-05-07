"""Sync constitution names and descriptions between JSON data and Excel.

Two modes:
    export  — read every JSON in src/data/constitutions/ and write a single
              XLSX containing one row per constitution (file, key, vi, en,
              passive_description_vi). Numeric/structural fields are NOT
              exported; only translatable text.
    import  — read the same XLSX back and patch the matching JSON entries
              by ``key``. Rows with empty ``key`` are skipped. Other JSON
              fields (rarity, stat_bonuses, materials, ...) are untouched.

Run from the repo root:
    python scripts/sync_constitution_text.py export
    python scripts/sync_constitution_text.py import

The XLSX path defaults to ``docs/constitutions.xlsx`` and can be overridden
with ``--xlsx <path>``.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parent.parent
CONST_DIR = REPO_ROOT / "src" / "data" / "constitutions"
DEFAULT_XLSX = REPO_ROOT / "docs" / "constitutions.xlsx"
SHEET_NAME = "Constitutions"

TEXT_FIELDS: tuple[str, ...] = ("vi", "en", "passive_description_vi")
HEADERS: tuple[str, ...] = ("file", "key", *TEXT_FIELDS)


def _load_json(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, list):
        raise ValueError(f"{path.name}: expected a JSON list, got {type(data).__name__}")
    return data


def _dump_json(path: Path, data: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
        fh.write("\n")


def export_to_xlsx(xlsx_path: Path) -> int:
    files = sorted(CONST_DIR.glob("*.json"))
    if not files:
        print(f"No JSON files found in {CONST_DIR}", file=sys.stderr)
        return 0

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(list(HEADERS))

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="FFD9E1F2")
    for col_idx, _ in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total = 0
    for path in files:
        entries = _load_json(path)
        for entry in entries:
            row = [
                path.name,
                entry.get("key", ""),
                entry.get("vi", "") or "",
                entry.get("en", "") or "",
                entry.get("passive_description_vi", "") or "",
            ]
            ws.append(row)
            total += 1

    column_widths = {
        "file": 20,
        "key": 38,
        "vi": 28,
        "en": 32,
        "passive_description_vi": 80,
    }
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = column_widths[header]

    last_col = get_column_letter(len(HEADERS))
    ws.auto_filter.ref = f"A1:{last_col}1"
    ws.freeze_panes = "A2"

    for row in ws.iter_rows(min_row=2, max_col=len(HEADERS)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"Exported {total} constitution(s) from {len(files)} file(s) -> {xlsx_path}")
    return total


def import_from_xlsx(xlsx_path: Path) -> int:
    if not xlsx_path.exists():
        print(f"XLSX not found: {xlsx_path}", file=sys.stderr)
        return 0

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print(f"{xlsx_path} is empty", file=sys.stderr)
        return 0

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    try:
        idx = {name: header.index(name) for name in HEADERS}
    except ValueError as exc:
        missing = [name for name in HEADERS if name not in header]
        raise SystemExit(f"XLSX missing required column(s): {missing}") from exc

    updates_by_file: dict[str, dict[str, dict[str, str]]] = {}
    skipped = 0
    for raw in rows[1:]:
        if raw is None or all(c is None or c == "" for c in raw):
            continue
        file_name = (raw[idx["file"]] or "").strip() if raw[idx["file"]] else ""
        key = (raw[idx["key"]] or "").strip() if raw[idx["key"]] else ""
        if not file_name or not key:
            skipped += 1
            continue
        text_patch = {
            field: ("" if raw[idx[field]] is None else str(raw[idx[field]]))
            for field in TEXT_FIELDS
        }
        updates_by_file.setdefault(file_name, {})[key] = text_patch

    if skipped:
        print(f"Skipped {skipped} row(s) with missing file/key.")

    total_changed = 0
    for file_name, key_to_patch in updates_by_file.items():
        path = CONST_DIR / file_name
        if not path.exists():
            print(f"  [skip] {file_name}: file not found", file=sys.stderr)
            continue

        entries = _load_json(path)
        file_changed = 0
        unmatched: list[str] = list(key_to_patch.keys())
        for entry in entries:
            key = entry.get("key")
            if not key or key not in key_to_patch:
                continue
            patch = key_to_patch[key]
            if key in unmatched:
                unmatched.remove(key)
            entry_changed = False
            for field, value in patch.items():
                if entry.get(field, "") != value:
                    entry[field] = value
                    entry_changed = True
            if entry_changed:
                file_changed += 1

        if file_changed:
            _dump_json(path, entries)
            print(f"  {file_name}: updated {file_changed} entry/entries")
        else:
            print(f"  {file_name}: no changes")

        for missing_key in unmatched:
            print(f"    [warn] key not found in {file_name}: {missing_key}", file=sys.stderr)

        total_changed += file_changed

    print(f"Imported updates: {total_changed} entry/entries changed.")
    return total_changed


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="mode", required=True)

    p_export = sub.add_parser("export", help="JSON -> XLSX")
    p_export.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX, help=f"Output XLSX path (default: {DEFAULT_XLSX})")

    p_import = sub.add_parser("import", help="XLSX -> JSON")
    p_import.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX, help=f"Input XLSX path (default: {DEFAULT_XLSX})")

    args = parser.parse_args(argv)

    if args.mode == "export":
        export_to_xlsx(args.xlsx)
    elif args.mode == "import":
        import_from_xlsx(args.xlsx)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
